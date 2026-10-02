import subprocess,sys,tempfile,unittest
from pathlib import Path
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches
ROOT=Path(__file__).resolve().parents[2]
class Embedded(unittest.TestCase):
 def test_standard_excel_backed_chart_has_actionable_refusal(self):
  with tempfile.TemporaryDirectory() as d:
   p=Path(d)/'chart.pptx';prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6]);data=CategoryChartData();data.categories=['Q1','Q2'];data.add_series('Sales',[10,20]);slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(6), Inches(4), data);prs.save(p)
   before=p.read_bytes();r=subprocess.run([sys.executable,str(ROOT/'src/python/import-model.py'),str(p),str(p)],capture_output=True,timeout=15)
   self.assertNotEqual(r.returncode,0);self.assertIn(b'embedded workbooks',r.stderr);self.assertIn(b'PNG, JPEG or GIF',r.stderr);self.assertEqual(p.read_bytes(),before)
if __name__=='__main__':unittest.main()
