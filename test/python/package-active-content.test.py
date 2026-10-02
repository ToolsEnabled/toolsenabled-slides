import io, json, os, subprocess, sys, tempfile, unittest, zipfile
from pathlib import Path
from pptx import Presentation
ROOT=Path(__file__).resolve().parents[2]
class ActiveContent(unittest.TestCase):
 def test_renamed_active_parts_are_refused_before_import_and_render(self):
  with tempfile.TemporaryDirectory() as directory:
   folder=Path(directory); good=folder/'good.pptx'; prs=Presentation();prs.slides.add_slide(prs.slide_layouts[6]);prs.save(good)
   with zipfile.ZipFile(good) as z: original={n:z.read(n) for n in z.namelist()}
   vectors={
    'cfb':{'ppt/media/innocent.dat':bytes.fromhex('d0cf11e0a1b11ae1')+b'\0'*256},
    'macro_type':{'[Content_Types].xml':original['[Content_Types].xml'].replace(b'presentationml.presentation.main+xml',b'presentation.macroEnabled.main+xml')},
    'ole_element':{'ppt/slides/slide1.xml':original['ppt/slides/slide1.xml'].replace(b'</p:spTree>',b'<p:oleObj/></p:spTree>')},
   }
   for kind in ['oleObject','package','vbaProject','vbaProjectSignature','control','activeXControl','activeXControlBinary']:
    vectors['rel_'+kind]={'ppt/_rels/presentation.xml.rels':original['ppt/_rels/presentation.xml.rels'].replace(b'</Relationships>',f'<Relationship Id="rUnsafe" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/{kind}" Target="media/innocent.bin"/></Relationships>'.encode()),'ppt/media/innocent.bin':b'benign test'}
   for label,changes in vectors.items():
    with self.subTest(label=label):
     deck=folder/(label+'.pptx')
     with zipfile.ZipFile(deck,'w',zipfile.ZIP_DEFLATED) as z:
      for n,b in (original|changes).items():z.writestr(n,b)
     result=subprocess.run([sys.executable,str(ROOT/'src/python/import-model.py'),str(deck),str(deck)],capture_output=True,timeout=15)
     self.assertNotEqual(result.returncode,0,'import accepted '+label)
     self.assertIn(b'Unsupported active content',result.stderr)
     # A persisted overlay from an earlier version must also be refused.
     model={'rev':0,'mode':'overlay','base':str(deck),'title':'Test','slides':[{'id':'s1','sourceIndex':0,'layout':'blank','elements':[],'notes':''}]}
     result=subprocess.run([sys.executable,str(ROOT/'src/python/render_pptx.py'),'-',str(folder/'out.pptx')],input=json.dumps(model).encode(),capture_output=True,timeout=15)
     self.assertNotEqual(result.returncode,0,'render accepted '+label)
     self.assertIn(b'Unsupported active content',result.stderr)
if __name__=='__main__':unittest.main()
