"""Portable content invariants using the generated synthetic overlay fixture."""
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "src" / "python"))
import copy, json, tempfile, unittest, zipfile
from pathlib import Path
from pptx import Presentation
from lxml import etree
import render_pptx as renderer
from import_pptx import derive_model
ROOT=Path(__file__).resolve().parents[2]
NS={'p':'http://schemas.openxmlformats.org/presentationml/2006/main'}
class CurrentContent(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp=tempfile.TemporaryDirectory(); cls.out=Path(cls.temp.name)/'rendered.pptx'
        cls.model=json.loads((ROOT/'data/model.json').read_text())
        renderer.render(copy.deepcopy(cls.model),str(cls.out))
        cls.base=Presentation(cls.model['base']); cls.deck=Presentation(cls.out)
    @classmethod
    def tearDownClass(cls): cls.temp.cleanup()
    def test_model_base_and_asset_graph(self):
        self.assertTrue(Path(self.model['base']).is_file())
        self.assertEqual(len(self.model['slides']),len(self.base.slides))
        for slide in self.model['slides']:
            self.assertEqual(len({e['id'] for e in slide['elements']}),len(slide['elements']))
            for e in slide['elements']:self.assertIn(e['ref']['sid'],[s.shape_id for s in self.base.slides[slide['src']].shapes])
    def test_stable_order(self):
        self.assertEqual([s['id'] for s in self.model['slides']],['s1','s2','s3'])
        self.assertEqual(len(self.deck.slides),3)
    def test_unchanged_text_and_rich_runs(self):
        for before,after in zip(self.base.slides,self.deck.slides):
            for a,b in zip(before.shapes,after.shapes):
                if a.has_text_frame:
                    self.assertEqual(a.text,b.text)
                    summary=lambda sh:[[(r.text,r.font.name,r.font.size,r.font.bold,r.font.italic) for r in p.runs] for p in sh.text_frame.paragraphs]
                    self.assertEqual(summary(a),summary(b))
    def test_geometry_is_preserved(self):
        for before,after in zip(self.base.slides,self.deck.slides):
            for a,b in zip(before.shapes,after.shapes):self.assertEqual((a.left,a.top,a.width,a.height),(b.left,b.top,b.width,b.height))
    def test_overlay_text_delta(self):
        m=copy.deepcopy(self.model); e=m['slides'][0]['elements'][0]; e['text']='A revised title'
        out=Path(self.temp.name)/'delta.pptx';renderer.render(m,str(out));deck=Presentation(out)
        shape=next(s for s in deck.slides[0].shapes if s.shape_id==e['ref']['sid'])
        self.assertEqual(shape.text,'A revised title')
        self.assertEqual([s.text for s in deck.slides[1].shapes if s.has_text_frame],[s.text for s in self.base.slides[1].shapes if s.has_text_frame])
    def test_silent_click_transition(self):
        m=copy.deepcopy(self.model); slide=m['slides'][0];slide['transition']={'effect':'fade','duration':.35}
        out=Path(self.temp.name)/'animated.pptx';renderer.render(m,str(out))
        with zipfile.ZipFile(out) as z:
            xml=etree.fromstring(z.read('ppt/slides/slide1.xml'))
            self.assertIsNotNone(xml.find('.//p:transition/p:fade',NS))
    def test_fonts_and_notes(self):
        for slide in self.deck.slides:
            if slide.has_notes_slide:self.assertFalse(slide.notes_slide.notes_text_frame.text.strip())
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for p in shape.text_frame.paragraphs:
                        for run in p.runs:
                            if run.text.strip(): self.assertGreaterEqual(run.font.size.pt,20);self.assertEqual(run.font.name,'DejaVu Sans')
    def test_imported_pictures_survive(self):
        for before,after in zip(self.base.slides,self.deck.slides):
            self.assertEqual([s.image.blob for s in before.shapes if hasattr(s,'image')],[s.image.blob for s in after.shapes if hasattr(s,'image')])
    def test_package_has_no_external_dependencies(self):
        with zipfile.ZipFile(self.out) as z:
            self.assertIsNone(z.testzip())
            for name in z.namelist():
                if name.endswith('.rels'):self.assertNotIn(b'TargetMode="External"',z.read(name))
if __name__=='__main__':unittest.main()
