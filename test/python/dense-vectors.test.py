import math, subprocess, sys, tempfile, unittest, zipfile
from pathlib import Path
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'src/python'))
import package_safety as safety

class DenseVectors(unittest.TestCase):
 def test_dense_vector_slide_in_two_hundred_slide_deck_imports_and_renders(self):
  with tempfile.TemporaryDirectory() as d:
   folder = Path(d); deck = folder / 'dense.pptx'; model = folder / 'model.json'; out = folder / 'out.pptx'
   prs = Presentation()
   for i in range(200):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(100000, 100000, 3000000, 500000).text = f'Review {i + 1}'
    if i == 0:
     for j in range(180):
      vertices = [(round(50000 * math.cos(k * math.tau / 60)), round(50000 * math.sin(k * math.tau / 60))) for k in range(60)]
      shape = slide.shapes.build_freeform(*vertices[0]).add_line_segments(vertices[1:]).convert_to_shape(150000 + j % 18 * 500000, 1000000 + j // 18 * 400000)
      shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor(35, 103, 161)
   prs.save(deck)
   with zipfile.ZipFile(deck) as z:
    xml = z.read('ppt/slides/slide1.xml'); nodes = list(etree.fromstring(xml).iter()); attributes = sum(len(n.attrib) for n in nodes)
   print('Dense slide XML bytes/elements/attributes:', len(xml), len(nodes), attributes, flush=True)
   self.assertGreater(attributes, 20000); self.assertLess(attributes, 50000)
   safety.validate_package(deck)
   imported = subprocess.run([sys.executable, str(ROOT / 'src/python/import-model.py'), str(deck), str(deck)], capture_output=True, timeout=30)
   self.assertEqual(imported.returncode, 0, imported.stderr.decode())
   model.write_bytes(imported.stdout)
   rendered = subprocess.run([sys.executable, str(ROOT / 'src/python/render_pptx.py'), str(model), str(out)], capture_output=True, timeout=30)
   self.assertEqual(rendered.returncode, 0, rendered.stderr.decode())
   result = Presentation(out); self.assertEqual(len(result.slides), 200); self.assertEqual(len(result.slides[0].shapes), 181)
if __name__ == '__main__': unittest.main()
