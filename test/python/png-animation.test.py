import io, struct, sys, tempfile, unittest, zipfile, zlib
from pathlib import Path
from PIL import Image
from pptx import Presentation
ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'src/python'))
from package_safety import package_digest, raster_bytes, validate_package
from static_render_copy import prepare_render_copy

def chunk(kind, data):
    return struct.pack('>I', len(data)) + kind + data + struct.pack('>I', zlib.crc32(kind + data))

def png(animation=False):
    out = io.BytesIO()
    Image.new('RGBA', (16, 12), 'red').save(out, format='PNG', save_all=animation,
        append_images=[Image.new('RGBA', (16, 12), 'blue')] if animation else [])
    return out.getvalue()

def add_chunks(data, extra):
    return data[:33] + extra + data[33:]

def deck(destination, data, disguised=False):
    p = Presentation()
    p.slides.add_slide(p.slide_layouts[6]).shapes.add_picture(io.BytesIO(png()), 0, 0)
    temp = io.BytesIO(); p.save(temp)
    with zipfile.ZipFile(temp) as source, zipfile.ZipFile(destination, 'w') as output:
        for item in source.infolist():
            value = source.read(item)
            if item.filename == 'ppt/media/image1.png': value = data
            name = item.filename
            if disguised:
                name = name.replace('image1.png', 'image1.gif')
                if name.endswith('.rels'): value = value.replace(b'image1.png', b'image1.gif')
            output.writestr(name, value)

class Animation(unittest.TestCase):
    def test_apng_and_disguised_apng_render_only_first_frame_without_changing_source(self):
        for disguised in (False, True):
            with self.subTest(disguised=disguised), tempfile.TemporaryDirectory() as directory:
                source = Path(directory) / 'source.pptx'; output = Path(directory) / 'render.pptx'
                data = png(True); deck(source, data, disguised); original = source.read_bytes()
                validate_package(source)
                prepare_render_copy(source, output, package_digest(source))
                self.assertEqual(source.read_bytes(), original)
                with zipfile.ZipFile(output) as result:
                    images = [result.read(n) for n in result.namelist() if '/media/' in n]
                self.assertEqual(len(images), 1)
                with Image.open(io.BytesIO(images[0])) as image:
                    self.assertEqual(image.n_frames, 1)
                    self.assertEqual(image.convert('RGB').getpixel((0, 0)), (255, 0, 0))
                for kind in (b'acTL', b'fcTL', b'fdAT', b'msOG'): self.assertNotIn(kind, images[0])

    def test_import_rejects_large_later_apng_frames_and_frame_count(self):
        for frames, size in ((2, 16000), (12000, 16)):
            data = png(True); parts = [data[:8]]; offset = 8; seen = 0
            while offset < len(data):
                length = int.from_bytes(data[offset:offset+4], 'big'); kind = data[offset+4:offset+8]
                value = data[offset+8:offset+8+length]
                if kind == b'acTL': value = struct.pack('>II', frames, 0)
                if kind == b'fcTL':
                    seen += 1
                    if seen == 2: value = value[:4] + struct.pack('>II', size, size) + value[12:]
                parts.append(chunk(kind, value)); offset += length + 12
            data = b''.join(parts)
            with self.subTest(frames=frames, size=size), self.assertRaises(ValueError): raster_bytes(data)

    def test_embedded_gif_in_msog_is_refused_at_import_and_stripped_for_office(self):
        gif = io.BytesIO(); Image.new('RGB', (16, 12), 'red').save(gif, format='GIF')
        data = add_chunks(png(), chunk(b'msOG', b'\0' * 4 + gif.getvalue()))
        with self.assertRaises(ValueError): raster_bytes(data)
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / 'source.pptx'; output = Path(directory) / 'render.pptx'
            deck(source, data); prepare_render_copy(source, output, package_digest(source))
            with zipfile.ZipFile(output) as result:
                self.assertTrue(all(b'msOG' not in result.read(n) for n in result.namelist() if '/media/' in n))

if __name__ == '__main__': unittest.main()
