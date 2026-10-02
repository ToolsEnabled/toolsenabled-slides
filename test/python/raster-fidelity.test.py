import io,sys,tempfile,unittest,zipfile
from pathlib import Path
from PIL import Image
from pptx import Presentation
ROOT=Path(__file__).resolve().parents[2];sys.path.insert(0,str(ROOT/'src/python'))
from package_safety import validate_package
class Fidelity(unittest.TestCase):
 def test_validated_jpeg_and_animated_gif_bytes_remain_identical(self):
  with tempfile.TemporaryDirectory() as d:
   for kind in ('JPEG','GIF'):
    with self.subTest(kind=kind):
     blob=io.BytesIO();im=Image.new('RGB',(32,32),'red')
     if kind=='JPEG':im.save(blob,format=kind,quality=98,comment=b'original metadata')
     else:im.save(blob,format=kind,save_all=True,append_images=[Image.new('RGB',(32,32),'blue')],duration=[80,120],loop=3)
     raw=blob.getvalue();p=Path(d)/(kind+'.pptx');prs=Presentation();prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_picture(io.BytesIO(raw),0,0);prs.save(p)
     before=p.read_bytes();validate_package(p,sanitize_images=True);self.assertEqual(p.read_bytes(),before)
     with zipfile.ZipFile(p) as z:
      media=[z.read(n) for n in z.namelist() if '/media/' in n];self.assertEqual(media,[raw])
 def test_pixel_cap_is_import_only(self):
  with tempfile.TemporaryDirectory() as d:
   blob=io.BytesIO();Image.new('RGB',(5100,5000),'red').save(blob,format='PNG');prs=Presentation();prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_picture(io.BytesIO(blob.getvalue()),0,0);p=Path(d)/'large.pptx';prs.save(p)
   with self.assertRaisesRegex(ValueError,'Unsupported image'):validate_package(p)
   validate_package(p,import_limits=False)
if __name__=='__main__':unittest.main()
