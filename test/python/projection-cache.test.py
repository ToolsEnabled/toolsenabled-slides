import errno, io, os, sys, tempfile, time, unittest, zipfile
from pathlib import Path
from unittest import mock
from PIL import Image
from pptx import Presentation
ROOT = Path(__file__).resolve().parents[2]; sys.path.insert(0, str(ROOT / 'src/python'))
from package_safety import package_digest
from static_render_copy import prepare_render_copy
import projection_cache as cache_module

class Cache(unittest.TestCase):
    def test_hash_cache_reuses_projection_and_rejects_corrupted_entries(self):
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory); source = folder / 'source.pptx'; cache = folder / 'cache'
            data = io.BytesIO(); Image.new('RGB', (16, 12), 'red').save(data, format='GIF', save_all=True,
                append_images=[Image.new('RGB', (16, 12), 'blue')])
            p = Presentation(); p.slides.add_slide(p.slide_layouts[6]).shapes.add_picture(io.BytesIO(data.getvalue()), 0, 0); p.save(source)
            digest = package_digest(source)
            prepare_render_copy(source, folder / 'one.pptx', digest, cache)
            with mock.patch('PIL.Image.open', side_effect=AssertionError('cached GIF must not decode again')):
                prepare_render_copy(source, folder / 'two.pptx', digest, cache)
            with zipfile.ZipFile(folder / 'one.pptx') as a, zipfile.ZipFile(folder / 'two.pptx') as b:
                self.assertEqual({n:a.read(n) for n in a.namelist()}, {n:b.read(n) for n in b.namelist()})
                self.assertTrue(all(a.getinfo(n).compress_type == zipfile.ZIP_STORED for n in a.namelist() if 'slides-first-frame' in n))
            cached = next(cache.glob('*/image.png')); cached.write_bytes(b'corrupted')
            original_open = Image.open
            with mock.patch('PIL.Image.open', wraps=original_open) as decoder:
                prepare_render_copy(source, folder / 'three.pptx', digest, cache)
                self.assertTrue(decoder.called)
            self.assertFalse(any(p.name.startswith('.pending-') for p in cache.iterdir()))

class CachePublication(unittest.TestCase):
    def test_publish_races_and_io_errors_do_not_fail_projection(self):
        for number in (errno.ENOTEMPTY, errno.EEXIST, errno.ENOSPC):
            with self.subTest(errno=number), tempfile.TemporaryDirectory() as folder:
                with mock.patch.object(Path, 'rename', side_effect=OSError(number, 'synthetic cache publish failure')):
                    cache_module.write(folder, 'a' * 64, b'projection')
                self.assertFalse(list(Path(folder).glob('.pending-*')))

    def test_stale_pending_directories_are_swept(self):
        with tempfile.TemporaryDirectory() as folder:
            stale = Path(folder) / '.pending-dead'; stale.mkdir(); (stale / 'image.png').write_bytes(b'abandoned')
            os.utime(stale, (time.time() - 3600,) * 2)
            cache_module.write(folder, 'a' * 64, b'projection')
            self.assertFalse(stale.exists())

if __name__ == '__main__': unittest.main()
