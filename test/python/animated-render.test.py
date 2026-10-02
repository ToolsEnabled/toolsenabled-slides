import io,json,struct,subprocess,sys,tempfile,time,unittest,zipfile
from pathlib import Path
from unittest import mock
from PIL import Image, GifImagePlugin
from pptx import Presentation
ROOT=Path(__file__).resolve().parents[2];sys.path.insert(0,str(ROOT/'src/python'))
import package_safety as safety
import render_pptx as renderer
from import_pptx import derive_model

def animation(color,frames=400):
 stream=io.BytesIO();Image.new('P',(1920,1080),color).save(stream,format='GIF');data=stream.getvalue();offset=13+(3*(2**((data[10]&7)+1)) if data[10]&128 else 0);start=data.index(b',',offset)
 return data[:start]+data[start:-1]*frames+b';'

class Animated(unittest.TestCase):
 def test_receipt_hashing_supports_python_310_without_file_digest(self):
  with tempfile.TemporaryDirectory() as d:
   source=Path(d)/'payload';source.write_bytes(b'package content' * 100000)
   expected=safety.hashlib.sha256(source.read_bytes()).hexdigest()
   with mock.patch.dict(vars(safety.hashlib)):
    vars(safety.hashlib).pop('file_digest',None)
    self.assertEqual(safety.package_digest(source),expected)
 def test_only_first_frame_is_decoded_and_truncated_later_frame_is_preserved(self):
  data=animation(1,3); original=GifImagePlugin.GifImageFile.load; seen=[]
  def counted(image,*args,**kwargs):
   seen.append(image.tell());return original(image,*args,**kwargs)
  with mock.patch.object(GifImagePlugin.GifImageFile,'load',counted):self.assertEqual(safety.raster_bytes(data),data)
  self.assertTrue(seen);self.assertEqual(set(seen),{0},'later frames must not be decoded')
  truncated=data[:-100];self.assertEqual(safety.raster_bytes(truncated),truncated)
 def test_four_400_frame_1080p_gifs_render_within_twenty_seconds(self):
  with tempfile.TemporaryDirectory() as d:
   folder=Path(d);deck=folder/'animation.pptx';prs=Presentation()
   for i in range(4):prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_picture(io.BytesIO(animation(i)+bytes([i])),0,0)
   prs.save(deck);model=folder/'model.json';out=folder/'out.pptx'
   model.write_text(json.dumps(derive_model(prs,'Animated',str(deck))))
   start=time.monotonic()
   rendered=subprocess.run([sys.executable,str(ROOT/'src/python/render_pptx.py'),str(model),str(out)],capture_output=True,timeout=20)
   self.assertEqual(rendered.returncode,0,rendered.stderr.decode());self.assertTrue(out.exists())
   with zipfile.ZipFile(deck) as before,zipfile.ZipFile(out) as after:
    self.assertEqual(sorted(before.read(n) for n in before.namelist() if '/media/' in n),sorted(after.read(n) for n in after.namelist() if '/media/' in n))
   imported=subprocess.run([sys.executable,str(ROOT/'src/python/import-model.py'),str(deck),str(deck)],capture_output=True,timeout=20)
   self.assertEqual(imported.returncode,0,imported.stderr.decode())
   with mock.patch.object(safety,'MAX_DECODE_PIXELS',100):
    with self.assertRaisesRegex(ValueError,'Unsupported image'):safety.validate_package(deck)
   print('Animated import + render seconds:',round(time.monotonic()-start,3),flush=True)
 def test_render_receipt_skips_validation_only_for_identical_bytes(self):
  with tempfile.TemporaryDirectory() as d:
   base=Path(d)/'base.pptx';out=Path(d)/'out.pptx';renderer.render({'slides':[{'id':'s1','elements':[]}]},base)
   model={'mode':'overlay','base':str(base),'slides':[{'id':'s1','src':0,'elements':[]}]}
   with mock.patch.object(safety,'validate_package',side_effect=AssertionError('unchanged input revalidated')), mock.patch.object(renderer,'validate_package',side_effect=AssertionError('unchanged input revalidated')):
    renderer.render_overlay(model,out)
   with zipfile.ZipFile(base,'a') as z:z.writestr('ppt/embeddings/evil.bin',b'bad')
   with self.assertRaisesRegex(ValueError,'Unsupported active content'):renderer.render_overlay(model,out)
if __name__=='__main__':unittest.main()
