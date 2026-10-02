#!/usr/bin/env python
"""Real PPTX regressions for first-deck typography and visible bullet glyphs."""
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "src" / "python"))
import tempfile
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

import render_pptx as renderer


class FirstDeckTests(unittest.TestCase):
    def test_set_bullets_content_renders_on_title_and_subtitle_elements(self):
        model = {"slides": [{"layout": "title", "elements": [
            {"id": "t", "type": "title", "items": ["Title item"]},
            {"id": "s", "type": "subtitle", "items": ["First item", "Second item"]},
        ]}]}
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "bullets.pptx"
            renderer.render_native(model, output)
            shapes = {s.name: s for s in Presentation(output).slides[0].shapes}
        for name, expected in [("t", ["Title item"]), ("s", ["First item", "Second item"])]:
            paras = shapes["Suite element " + name].text_frame.paragraphs
            self.assertEqual([p.text for p in paras], expected)
            self.assertTrue(all(p._p.find("./" + qn("a:pPr") + "/" + qn("a:buChar")) is not None for p in paras))

    def test_bullets_are_explicit_and_retain_run_formatting(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(4))
        shape.text_frame.text = "Previous text"
        run = shape.text_frame.paragraphs[0].runs[0]
        run.font.size = Pt(23)
        run.font.name = "Arial"
        # A source textbox can explicitly suppress bullets. set_bullets must
        # replace that suppression, including on newly cloned paragraphs.
        ppr = shape.text_frame.paragraphs[0]._p.get_or_add_pPr()
        ppr.append(OxmlElement("a:buNone"))
        renderer.apply_bullets(shape.text_frame, ["One shared goal", "A measurable result"])
        for para in shape.text_frame.paragraphs:
            ppr = para._p.find(qn("a:pPr"))
            self.assertIsNotNone(ppr)
            self.assertIsNone(ppr.find(qn("a:buNone")))
            bullet = ppr.find(qn("a:buChar"))
            self.assertIsNotNone(bullet, "Every bullet paragraph needs a visible DrawingML glyph")
            self.assertEqual(bullet.get("char"), "•")
            self.assertEqual(para.runs[0].font.size.pt, 23)
            self.assertEqual(para.runs[0].font.name, "Arial")

    def test_native_default_is_styled_and_explicit_overrides_survive(self):
        model = {"slides": [
            {"id": "s1", "layout": "title", "elements": [
                {"id": "t", "type": "title", "text": "A clear plan"},
                {"id": "s", "type": "subtitle", "text": "Evidence. Decisions. Next steps."}]},
            {"id": "s2", "layout": "content", "elements": [
                {"id": "t2", "type": "title", "text": "What we learned"},
                {"id": "b", "type": "body", "items": ["One shared goal", "A measurable result"]},
                {"id": "custom", "type": "heading", "text": "Custom heading",
                 "style": {"font": "Courier New", "size": 19, "color": "#112233"},
                 "box": {"x": .1, "y": .8, "w": .4, "h": .1}}]}]}
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "default.pptx"
            renderer.render_native(model, output)
            prs = Presentation(output)
        title = prs.slides[0].shapes.title.text_frame.paragraphs[0].runs[0]
        self.assertEqual(title.font.name, "Arial")
        self.assertGreaterEqual(title.font.size.pt, 36)
        self.assertTrue(title.font.bold)
        self.assertEqual(str(title.font.color.rgb), "0F172A")
        shapes = {s.name: s for s in prs.slides[1].shapes}
        body = shapes["Suite element b"]
        for para in body.text_frame.paragraphs:
            self.assertEqual(para.runs[0].font.name, "Arial")
            self.assertEqual(para.runs[0].font.size.pt, 24)
            self.assertGreaterEqual(para.space_after.pt, 10)
            self.assertIsNotNone(para._p.find("./" + qn("a:pPr") + "/" + qn("a:buChar")))
        custom = shapes["Suite element custom"].text_frame.paragraphs[0].runs[0]
        self.assertEqual(custom.font.name, "Courier New")
        self.assertEqual(custom.font.size.pt, 19)
        self.assertEqual(str(custom.font.color.rgb), "112233")


if __name__ == "__main__":
    unittest.main()
