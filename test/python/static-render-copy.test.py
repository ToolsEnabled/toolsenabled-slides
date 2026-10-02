import io, sys, tempfile, unittest, zipfile
from pathlib import Path
from unittest import mock
from PIL import Image, GifImagePlugin
from pptx import Presentation
ROOT=Path(__file__).resolve().parents[2];sys.path.insert(0,str(ROOT/'src/python'))
from static_render_copy import prepare_render_copy
from package_safety import package_digest, validate_package

class StaticCopy(unittest.TestCase):
 def test_first_frame_alpha_and_dimensions_without_decoding_later_frames(self):
  with tempfile.TemporaryDirectory() as d:
   source=Path(d)/'source.pptx';output=Path(d)/'render.pptx';data=io.BytesIO()
   first=Image.new('RGBA',(16,12),(0,0,0,0));first.putpixel((3,4),(255,0,0,255))
   first.save(data,format='GIF',save_all=True,append_images=[Image.new('RGBA',(16,12),(0,255,0,255))],disposal=2)
   p=Presentation();p.slides.add_slide(p.slide_layouts[6]).shapes.add_picture(io.BytesIO(data.getvalue()),0,0);p.save(source);original=source.read_bytes();seen=[];load=GifImagePlugin.GifImageFile.load
   def tracked(image,*args,**kwargs):seen.append(image.tell());return load(image,*args,**kwargs)
   with mock.patch.object(GifImagePlugin.GifImageFile,'load',tracked):receipt=prepare_render_copy(source,output,package_digest(source))
   self.assertEqual(set(seen),{0});self.assertEqual(receipt['gifPartsReplaced'],1);self.assertEqual(source.read_bytes(),original)
   image=Image.open(io.BytesIO(Presentation(output).slides[0].shapes[0].image.blob));self.assertEqual(image.format,'PNG');self.assertEqual(image.size,(16,12));self.assertEqual(image.getpixel((3,4)),(255,0,0,255));self.assertEqual(image.getpixel((0,0))[3],0)
   validate_package(output,import_limits=False)
 def test_hash_mismatch_and_same_destination_refuse_without_changes(self):
  with tempfile.TemporaryDirectory() as d:
   source=Path(d)/'source.pptx';output=Path(d)/'render.pptx';Presentation().save(source);original=source.read_bytes()
   with self.assertRaisesRegex(ValueError,'hash changed'):prepare_render_copy(source,output,'0'*64)
   self.assertFalse(output.exists())
   with self.assertRaisesRegex(ValueError,'separate private'):prepare_render_copy(source,source,package_digest(source))
   self.assertEqual(source.read_bytes(),original)
 def test_static_media_needs_no_projection_and_existing_destination_is_preserved(self):
  with tempfile.TemporaryDirectory() as d:
   source=Path(d)/'source.pptx';output=Path(d)/'render.pptx';data=io.BytesIO();Image.new('RGB',(12,10),'blue').save(data,format='GIF');p=Presentation();p.slides.add_slide(p.slide_layouts[6]).shapes.add_picture(io.BytesIO(data.getvalue()),0,0);p.save(source)
   receipt=prepare_render_copy(source,output,package_digest(source));self.assertEqual(receipt['gifPartsReplaced'],0);self.assertEqual(source.read_bytes(),output.read_bytes())
   output.write_bytes(b'keep existing destination')
   with self.assertRaises(FileExistsError):prepare_render_copy(source,output,package_digest(source))
   self.assertEqual(output.read_bytes(),b'keep existing destination')
if __name__=='__main__':unittest.main()
