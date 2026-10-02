"""Reusable final artifact acceptance; all content is synthetic."""
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "src" / "python"))
import json, hashlib, unittest, zipfile
from pathlib import Path
from pptx import Presentation
ROOT=Path(__file__).resolve().parents[2]
class FinalDeckAcceptance(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.model=json.loads((ROOT/'data/model.json').read_text());cls.file=ROOT.parent/'presentation.pptx';cls.deck=Presentation(cls.file)
    def test_slide_count_and_order(self):
        self.assertEqual(len(self.deck.slides),3)
        self.assertEqual([s['id'] for s in self.model['slides']],['s1','s2','s3'])
    def test_expected_titles(self):
        expected=['Project overview','Evidence and decisions','Next steps']
        for slide,title in zip(self.deck.slides,expected):self.assertIn(title,[s.text for s in slide.shapes if s.has_text_frame])
    def test_every_object_is_inside_slide(self):
        for slide in self.deck.slides:
            for shape in slide.shapes:
                self.assertGreaterEqual(shape.left,0);self.assertGreaterEqual(shape.top,0)
                self.assertLessEqual(shape.left+shape.width,self.deck.slide_width);self.assertLessEqual(shape.top+shape.height,self.deck.slide_height)
    def test_readable_text(self):
        for slide in self.deck.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for p in shape.text_frame.paragraphs:
                        for run in p.runs:
                            if run.text:self.assertGreaterEqual(run.font.size.pt,20)
    def test_notes_free(self):
        for slide in self.deck.slides:
            if slide.has_notes_slide:self.assertFalse(slide.notes_slide.notes_text_frame.text.strip())
    def test_footer_clearance(self):
        for slide in self.deck.slides:
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text:self.assertLess(shape.top+shape.height,self.deck.slide_height*.9)
    def test_render_marker_matches_bytes(self):
        marker=json.loads((ROOT/'data/render-state.json').read_text());raw=self.file.read_bytes()
        self.assertEqual(marker['rev'],self.model['rev']);self.assertEqual(marker['size'],len(raw));self.assertEqual(marker['sha256'],hashlib.sha256(raw).hexdigest())
    def test_valid_self_contained_package(self):
        with zipfile.ZipFile(self.file) as z:
            self.assertIsNone(z.testzip());self.assertIn('ppt/presentation.xml',z.namelist())
            self.assertFalse(any('vbaProject' in n or '/embeddings/' in n for n in z.namelist()))
if __name__=='__main__':unittest.main()
