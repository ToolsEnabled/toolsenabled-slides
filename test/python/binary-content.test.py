"""Every OPC part is classified by bytes, including disguised diagram/image targets."""
import io, struct, sys, tempfile, unittest, zipfile
from pathlib import Path
from PIL import Image
from pptx import Presentation
ROOT = Path(__file__).resolve().parents[2]; sys.path.insert(0, str(ROOT / 'src/python'))
from package_safety import validate_package

def package(path, data, relation='font', blip=False):
    deck = Presentation(); deck.slides.add_slide(deck.slide_layouts[6]); original = io.BytesIO(); deck.save(original)
    with zipfile.ZipFile(original) as source, zipfile.ZipFile(path, 'w') as result:
        for name in source.namelist():
            value = source.read(name)
            if name == '[Content_Types].xml':
                value = value.replace(b'</Types>', b'<Override PartName="/ppt/fonts/font1.fntdata" ContentType="application/x-fontdata"/></Types>')
            if name == 'ppt/slides/_rels/slide1.xml.rels':
                value = value.replace(b'</Relationships>', f'<Relationship Id="rIdF" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/{relation}" Target="../fonts/font1.fntdata"/></Relationships>'.encode())
            if name == 'ppt/slides/slide1.xml':
                node = b'<a:blip r:embed="rIdF"/>' if blip else b'<dgm:relIds xmlns:dgm="http://schemas.openxmlformats.org/drawingml/2006/diagram" r:dm="rIdF"/>'
                value = value.replace(b'</p:spTree>', b'<p:graphicFrame><a:graphic><a:graphicData>' + node + b'</a:graphicData></a:graphic></p:graphicFrame></p:spTree>')
            result.writestr(name, value)
        result.writestr('ppt/fonts/font1.fntdata', data)

class BinaryContent(unittest.TestCase):
    def test_font_xml_is_never_exempt(self):
        cases = [b'<r/>', b'<r>' + b'<n a="1234567890"/>' * 200000 + b'</r>', b'<!DOCTYPE r [<!ENTITY a "x"><!ENTITY b "&a;&a;">]><r>&b;</r>']
        with tempfile.TemporaryDirectory() as folder:
            for data in cases:
                with self.subTest(length=len(data)):
                    target = Path(folder) / 'deck.pptx'; package(target, data)
                    with self.assertRaises(ValueError): validate_package(target)

    def test_unsupported_rasters_cannot_hide_outside_media(self):
        values = []
        for fmt in ['BMP', 'TIFF', 'WEBP']:
            data = io.BytesIO(); Image.new('RGB', (12, 12)).save(data, format=fmt); values.append(data.getvalue())
        huge = bytearray(values[0]); struct.pack_into('<ii', huge, 18, 30000, 30000); values.append(bytes(huge))
        with tempfile.TemporaryDirectory() as folder:
            for relation in ['font', 'custom', 'image']:
                for data in values:
                    with self.subTest(relation=relation, magic=data[:4]):
                        target = Path(folder) / 'deck.pptx'; package(target, data, relation, True)
                        with self.assertRaises(ValueError): validate_package(target)

    def test_blip_must_resolve_to_validated_image(self):
        with tempfile.TemporaryDirectory() as folder:
            target = Path(folder) / 'deck.pptx'; package(target, b'<root/>', 'custom', True)
            with self.assertRaises(ValueError): validate_package(target)

    def test_svg_cannot_hide_as_custom_xml_outside_media(self):
        with tempfile.TemporaryDirectory() as folder:
            target = Path(folder) / 'deck.pptx'
            package(target, b'<svg xmlns="http://www.w3.org/2000/svg" width="30000" height="30000"/>', 'custom')
            with self.assertRaisesRegex(ValueError, 'Unsupported image'): validate_package(target)

    def test_mac_printer_records_do_not_hide_custom_doctypes(self):
        from package_safety import binary_kind
        payload = b'<!DOCTYPE plist [<!ENTITY a "x">]><plist>&a;</plist>'
        wrapped = struct.pack('<III', 2, 1, len(payload)) + payload + struct.pack('<II', 2, len(payload)) + payload
        self.assertIsNone(binary_kind(wrapped, 'printer.bin'))
        payload = b'<?xml version="1.0"?>' + payload
        wrapped = struct.pack('<III', 2, 1, len(payload)) + payload + struct.pack('<II', 2, len(payload)) + payload
        with self.assertRaises(ValueError): binary_kind(wrapped, 'printer.bin')

if __name__ == '__main__': unittest.main()
