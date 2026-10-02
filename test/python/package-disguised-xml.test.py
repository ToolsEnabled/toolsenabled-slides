import sys,tempfile,unittest,zipfile
from pathlib import Path
from pptx import Presentation
ROOT=Path(__file__).resolve().parents[2];sys.path.insert(0,str(ROOT/'src/python'))
from package_safety import validate_package,MAX_PART_ATTRIBUTES
class Disguised(unittest.TestCase):
 def test_case_variant_override_cannot_bypass_xml_budget(self):
  with tempfile.TemporaryDirectory() as d:
   p=Path(d)/'deck.pptx';prs=Presentation();prs.slides.add_slide(prs.slide_layouts[6]);prs.save(p)
   with zipfile.ZipFile(p) as z:parts={n:z.read(n) for n in z.namelist()}
   parts['[Content_Types].xml']=parts['[Content_Types].xml'].replace(b'</Types>',b'<Override PartName="/PPT/PAYLOAD.DAT" ContentType="application/vnd.openxmlformats-officedocument.presentationml.slide+xml"/></Types>')
   for encoding in ('utf-8','utf-16'):
    with self.subTest(encoding=encoding):
     payload=('<root '+' '.join(f'a{i}="x"' for i in range(MAX_PART_ATTRIBUTES+1))+'/>').encode(encoding)
     with zipfile.ZipFile(p,'w') as z:
      for n,b in parts.items():z.writestr(n,b)
      z.writestr('ppt/payload.dat',payload)
     with self.assertRaisesRegex(ValueError,'budget'):validate_package(p)
 def test_xml_without_declaration_or_xml_content_type_is_preflighted(self):
  with tempfile.TemporaryDirectory() as d:
   p=Path(d)/'deck.pptx';prs=Presentation();prs.slides.add_slide(prs.slide_layouts[6]);prs.save(p)
   with zipfile.ZipFile(p,'a') as z:z.writestr('ppt/payload.bin',b'<root>'+b'<n/>'*50001+b'</root>')
   with self.assertRaisesRegex(ValueError,'budget'):validate_package(p)
if __name__=='__main__':unittest.main()
