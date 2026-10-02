import io,sys,tempfile,unittest,zipfile
from pathlib import Path
from pptx import Presentation
ROOT=Path(__file__).resolve().parents[2];sys.path.insert(0,str(ROOT/'src/python'))
from package_safety import validate_package
class XMLSniff(unittest.TestCase):
 def test_ebcdic_and_unusual_ucs4_byte_orders_do_not_bypass_xml_checks(self):
  text='<?xml version="1.0" encoding="IBM037"?><!DOCTYPE root><root/>'
  variants=[text.encode('cp037'),text.replace('IBM037','IBM500').encode('cp500')]
  data='<?xml version="1.0"?><!DOCTYPE root><root/>'.encode('utf-32-be')
  for order in ((1,0,3,2),(2,3,0,1)):
   variants.append(b''.join(bytes(block[i] for i in order) for block in (data[n:n+4] for n in range(0,len(data),4))))
  with tempfile.TemporaryDirectory() as d:
   for index,data in enumerate(variants):
    with self.subTest(index=index):
     deck=Path(d)/'deck.pptx';Presentation().save(deck)
     with zipfile.ZipFile(deck,'a') as z:z.writestr('ppt/theme/payload.bin',data)
     with self.assertRaises(ValueError):validate_package(deck)
 def test_whole_disguised_parts_are_counted_for_all_xml_encodings(self):
  with tempfile.TemporaryDirectory() as d:
   deck=Path(d)/'deck.pptx';prs=Presentation();prs.slides.add_slide(prs.slide_layouts[6]);prs.save(deck)
   with zipfile.ZipFile(deck) as z:parts={n:z.read(n) for n in z.namelist()}
   body='<root>'+('<n/>'*50001)+'</root>'
   # The unknown content type keeps this test focused on the sniff fallback.
   variants={'whitespace':(' '*20000+body).encode(),'split-multibyte':('<?xml version="1.0" encoding="UTF-8"?><root><!--'+'x'*(255-len('<?xml version="1.0" encoding="UTF-8"?><root><!--'))+'é-->'+('<n/>'*50001)+'</root>').encode()}
   for encoding in ('utf-16-be','utf-16-le','utf-32-be','utf-32-le'):variants[encoding]=(' '*300+body).encode(encoding)
   for label,data in variants.items():
    with self.subTest(label=label):
     with zipfile.ZipFile(deck,'w') as z:
      for n,b in parts.items():z.writestr(n,b)
      z.writestr('ppt/theme/payload.bin',data)
     with self.assertRaisesRegex(ValueError,'budget'):validate_package(deck)
if __name__=='__main__':unittest.main()
