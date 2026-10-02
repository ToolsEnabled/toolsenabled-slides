import io, subprocess, sys, tempfile, unittest, zipfile
from pathlib import Path
from pptx import Presentation
from PIL import Image
ROOT=Path(__file__).resolve().parents[2]
class Images(unittest.TestCase):
 def test_disguised_svg_and_other_non_raster_image_parts_are_refused(self):
  with tempfile.TemporaryDirectory() as directory:
   folder=Path(directory); png=folder/'good.png';Image.new('RGB',(8,8),'red').save(png);prs=Presentation();prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_picture(str(png),0,0);good=folder/'good.pptx';prs.save(good)
   with zipfile.ZipFile(good) as z: original={n:z.read(n) for n in z.namelist()}
   svg=b'<svg xmlns="http://www.w3.org/2000/svg" xmlns:xlink="http://www.w3.org/1999/xlink"><image xlink:href="file:///not-read/secret.png"/></svg>'
   for label,data in [('SVG',svg),('PDF',b'%PDF-1.7'),('WMF',bytes.fromhex('d7cdc69a')+b'\0'*32)]:
    with self.subTest(label=label):
     deck=folder/(label+'.pptx')
     with zipfile.ZipFile(deck,'w',zipfile.ZIP_DEFLATED) as z:
      for n,b in original.items():z.writestr(n,data if n=='ppt/media/image1.png' else b)
     result=subprocess.run([sys.executable,str(ROOT/'src/python/import-model.py'),str(deck),str(deck)],capture_output=True,timeout=15)
     self.assertNotEqual(result.returncode,0,'accepted disguised '+label)
     self.assertIn(b'Unsupported image',result.stderr)
 def test_image_relationship_cannot_hide_svg_under_a_non_media_binary_part(self):
  with tempfile.TemporaryDirectory() as directory:
   folder=Path(directory);png=folder/'good.png';Image.new('RGB',(8,8),'red').save(png);prs=Presentation();prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_picture(str(png),0,0);deck=folder/'hidden.pptx';prs.save(deck)
   with zipfile.ZipFile(deck) as z:contents={n:z.read(n) for n in z.namelist()}
   contents.pop('ppt/media/image1.png')
   contents['ppt/other/payload.dat']=b'<svg xmlns="http://www.w3.org/2000/svg"><image href="file:///not-read/secret.png"/></svg>'
   contents['[Content_Types].xml']=contents['[Content_Types].xml'].replace(b'</Types>',b'<Default Extension="dat" ContentType="application/octet-stream"/></Types>')
   contents['ppt/slides/_rels/slide1.xml.rels']=contents['ppt/slides/_rels/slide1.xml.rels'].replace(b'../media/image1.png',b'../other/payload.dat')
   with zipfile.ZipFile(deck,'w') as z:
    for n,b in contents.items():z.writestr(n,b)
   # Test the package boundary directly: Office may decode parts python-pptx leaves opaque.
   sys.path.insert(0,str(ROOT/'src/python'))
   from package_safety import validate_package
   with self.assertRaisesRegex(ValueError,'Unsupported image'):validate_package(deck)
 def test_import_preserves_validated_raster_bytes_and_pixels(self):
  with tempfile.TemporaryDirectory() as directory:
   folder=Path(directory);png=folder/'good.png';Image.new('RGB',(8,8),'red').save(png);prs=Presentation();prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_picture(str(png),0,0);deck=folder/'good.pptx';prs.save(deck)
   with zipfile.ZipFile(deck) as z: contents={n:z.read(n) for n in z.namelist()}
   contents['ppt/media/image1.png']+=b'UNTRUSTED_TRAILING_METADATA'
   with zipfile.ZipFile(deck,'w',zipfile.ZIP_DEFLATED) as z:
    for n,b in contents.items():z.writestr(n,b)
   result=subprocess.run([sys.executable,str(ROOT/'src/python/import-model.py'),str(deck),str(deck)],capture_output=True,timeout=15)
   self.assertEqual(result.returncode,0,result.stderr)
   with zipfile.ZipFile(deck) as z:
    data=z.read('ppt/media/image1.png');self.assertEqual(data,contents['ppt/media/image1.png']);self.assertEqual(Image.open(io.BytesIO(data)).convert('RGB').getpixel((0,0)),(255,0,0))
if __name__=='__main__':unittest.main()
