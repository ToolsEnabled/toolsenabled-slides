import sys,tempfile,unittest,zipfile
from pathlib import Path
from pptx import Presentation
ROOT=Path(__file__).resolve().parents[2];sys.path.insert(0,str(ROOT/'src/python'))
from package_safety import validate_package
class Comments(unittest.TestCase):
 def test_comments_in_content_types_relationships_and_slides_are_inert(self):
  with tempfile.TemporaryDirectory() as d:
   p=Path(d)/'deck.pptx';prs=Presentation();prs.slides.add_slide(prs.slide_layouts[6]);prs.save(p)
   with zipfile.ZipFile(p) as z:parts={n:z.read(n) for n in z.namelist()}
   for n,ending in [('[Content_Types].xml',b'</Types>'),('ppt/slides/_rels/slide1.xml.rels',b'</Relationships>'),('ppt/slides/slide1.xml',b'</p:sld>')]:parts[n]=parts[n].replace(ending,b'<!-- inert comment -->'+ending)
   with zipfile.ZipFile(p,'w') as z:
    for n,b in parts.items():z.writestr(n,b)
   validate_package(p)
   parts['ppt/slides/_rels/slide1.xml.rels']=parts['ppt/slides/_rels/slide1.xml.rels'].replace(b'</Relationships>',b'<Relationship Id="bad" Type="http://example/oleObject" Target="bad.bin"/></Relationships>')
   with zipfile.ZipFile(p,'w') as z:
    for n,b in parts.items():z.writestr(n,b)
   with self.assertRaisesRegex(ValueError,'Unsupported active content'):validate_package(p)
if __name__=='__main__':unittest.main()
