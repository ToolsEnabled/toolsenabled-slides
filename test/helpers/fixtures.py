"""Synthetic imported deck for the original integration suite; no user content."""
import sys, json, shutil, hashlib, base64
from pathlib import Path
root = Path(sys.argv[1]).resolve()
sys.path.insert(0, str(root / "src" / "python"))
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from PIL import Image
from import_pptx import derive_model

data = root / 'data'; data.mkdir(parents=True, exist_ok=True)
assets = root.parent / 'assets'; assets.mkdir(exist_ok=True)
(assets/'fixture.mp4').write_bytes(base64.b64decode((root/'test/fixtures/h264-blue-32x32.mp4.base64').read_text()))
png = assets / 'fixture.png'; Image.new('RGB', (32, 32), '#2563eb').save(png)
prs = Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
for i in range(3):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(.6), Inches(.4), Inches(11), Inches(.7))
    title.text = ['Project overview', 'Evidence and decisions', 'Next steps'][i]
    box=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(.6), Inches(1.5), Inches(6), Inches(2))
    box.fill.solid(); box.fill.fore_color.rgb=RGBColor.from_string('E2E8F0')
    box.text = 'Editable rounded card'
    decor=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8), Inches(1.5), Inches(3), Inches(2))
    decor.fill.solid(); decor.fill.fore_color.rgb=RGBColor.from_string('2563EB')
    slide.shapes.add_picture(str(png), Inches(11.5), Inches(5.5), Inches(.7), Inches(.7))
    body=slide.shapes.add_textbox(Inches(.6), Inches(4), Inches(10), Inches(1.5))
    body.text='Clear evidence\nA measurable next step'
    for shape in slide.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                for run in p.runs: run.font.name='DejaVu Sans'; run.font.size=Pt(24)
base=data/'base.pptx'; prs.save(base)
model=derive_model(prs, 'Sample presentation.pptx', str(base))
for slide in model['slides']:
    for i,d in enumerate(slide['decor']): d['id']=f"{slide['id']}_d{i}"
(data/'model.json').write_text(json.dumps(model))
(data/'element-history.json').write_text(json.dumps({'schemaVersion':1,'modelRev':0,'entries':{}}))
(data/'state.json').write_text(json.dumps({'paused':False,'pausedBy':None,'pausedAt':None,'pdfRev':None}))
(data/'board.json').write_text(json.dumps({'notes':[],'tasks':[],'loops':[]}))
for dest in (root.parent/'presentation.pptx', data/'thumbs-src.pptx'): shutil.copyfile(base,dest)
raw=base.read_bytes()
(data/'render-state.json').write_text(json.dumps({'rev':0,'size':len(raw),'sha256':hashlib.sha256(raw).hexdigest(),'renderedAt':'2026-01-01T00:00:00.000Z'}))
